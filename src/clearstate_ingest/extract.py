"""Readers. One job: turn a source into a header plus rows, and nothing else.

Interpretation happens later, so that a parse can always be replayed from the
archived original rather than re-fetched from a source that has since moved on.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence

from .contracts import Contract
from .schema import Schema, infer_schema, normalise


@dataclass
class SourceData:
    """Raw, uninterpreted content of one source read."""

    source_id: str
    header: List[str]
    rows: List[List[Any]]
    origin: str = ""
    content_hash: str = ""

    def __len__(self) -> int:
        return len(self.rows)

    @property
    def normalised_header(self) -> List[str]:
        return [normalise(h) for h in self.header]

    def as_dicts(self) -> List[Dict[str, Any]]:
        """Rows keyed by normalised source column name."""
        keys = self.normalised_header
        out: List[Dict[str, Any]] = []
        for row in self.rows:
            out.append({k: (row[i] if i < len(row) else None) for i, k in enumerate(keys)})
        return out

    def infer(self, sample: int = 500) -> Schema:
        return infer_schema(self.normalised_header, self.rows[:sample])


def _hash_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def read_csv(path: str, header_row: int = 1, delimiter: str = ",", encoding: str = "utf-8-sig") -> SourceData:
    """Read a CSV, skipping ``header_row - 1`` preamble lines.

    Real exports carry a title line and a blank line above the header; that is
    what ``header_row`` exists for.
    """
    with io.open(path, "r", encoding=encoding, newline="") as fh:
        raw = fh.read()
    reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
    all_rows = [r for r in reader]
    if not all_rows:
        raise ValueError("{} is empty".format(path))
    idx = max(0, header_row - 1)
    if idx >= len(all_rows):
        raise ValueError("header_row {} is past the end of {}".format(header_row, path))
    header = [c for c in all_rows[idx]]
    body = [r for r in all_rows[idx + 1 :] if any(str(c).strip() for c in r)]
    return SourceData(
        source_id=os.path.basename(path),
        header=header,
        rows=body,
        origin=path,
        content_hash=_hash_bytes(raw.encode("utf-8")),
    )


def read_excel(path: str, sheet_name: Optional[str] = None, header_row: int = 1) -> SourceData:
    """Read one worksheet via openpyxl. Imported lazily so the core has no deps."""
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError(
            "reading .xlsx needs openpyxl - pip install openpyxl, or export the tab to CSV"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not grid:
        raise ValueError("{} has no rows".format(path))
    idx = max(0, header_row - 1)
    header = [("" if c is None else str(c)) for c in grid[idx]]
    body = [list(r) for r in grid[idx + 1 :] if any(c is not None and str(c).strip() for c in r)]
    with open(path, "rb") as fh:
        content_hash = _hash_bytes(fh.read())
    return SourceData(
        source_id=os.path.basename(path),
        header=header,
        rows=body,
        origin="{}#{}".format(path, ws.title),
        content_hash=content_hash,
    )


def read_google_sheet(spreadsheet_id: str, tab: str, header_row: int = 1) -> SourceData:
    """Read a tab through the Sheets API.

    Kept behind a lazy import so the package installs and tests without any
    Google dependency. In production this runs inside the Cloud Run extractor
    with a service account that has read access to the sheet only.
    """
    try:  # pragma: no cover - requires network + credentials
        from googleapiclient.discovery import build  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "reading Google Sheets needs google-api-python-client "
            "(pip install -e '.[gcp]')"
        ) from exc

    service = build("sheets", "v4")  # pragma: no cover
    result = (  # pragma: no cover
        service.spreadsheets()
        .values()
        .get(spreadsheetId=spreadsheet_id, range=tab)
        .execute()
    )
    values = result.get("values", [])  # pragma: no cover
    if not values:  # pragma: no cover
        raise ValueError("tab {!r} is empty".format(tab))
    idx = max(0, header_row - 1)  # pragma: no cover
    return SourceData(  # pragma: no cover
        source_id=spreadsheet_id,
        header=list(values[idx]),
        rows=[list(r) for r in values[idx + 1 :]],
        origin="sheets://{}/{}".format(spreadsheet_id, tab),
        content_hash=_hash_bytes(repr(values).encode("utf-8")),
    )


def read_source(contract: Contract, base_dir: str = ".") -> SourceData:
    """Dispatch to the right reader using the contract's location settings."""
    loc = dict(contract.location)
    header_row = int(loc.get("header_row", 1))
    kind = contract.kind.lower()

    if kind in ("csv", "sheet_export"):
        path = loc.get("path")
        if not path:
            raise ValueError("contract '{}' has no 'path'".format(contract.source_id))
        if not os.path.isabs(path):
            path = os.path.join(base_dir, path)
        data = read_csv(
            path,
            header_row=header_row,
            delimiter=loc.get("delimiter", ","),
            encoding=loc.get("encoding", "utf-8-sig"),
        )
    elif kind in ("excel", "xlsx"):
        path = loc.get("path")
        if not path:
            raise ValueError("contract '{}' has no 'path'".format(contract.source_id))
        if not os.path.isabs(path):
            path = os.path.join(base_dir, path)
        data = read_excel(path, sheet_name=loc.get("sheet_name"), header_row=header_row)
    elif kind == "google_sheet":
        data = read_google_sheet(  # pragma: no cover
            spreadsheet_id=loc["spreadsheet_id"], tab=loc["tab"], header_row=header_row
        )
    else:
        raise ValueError("unsupported source kind {!r}".format(contract.kind))

    data.source_id = contract.source_id
    return data
